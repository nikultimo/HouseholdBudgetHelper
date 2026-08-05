from __future__ import annotations
from datetime import date
from typing import Any
import openpyxl
from openpyxl.formula.translate import Translator

from salary.calculator import calculate_payment

from excel.constants import (
    LEDGER_COL_DELTA,
    LEDGER_COL_ENTRY_ID,
    LEDGER_COL_KIND,
    LEDGER_COL_POSTED_AT,
    LEDGER_COL_REVERSES_ID,
    SHEET_TRANSACTIONS,
    SHEET_CAPITAL,
    SHEET_SETTINGS,
)
from excel.ledger import (
    ensure_ledger_schema,
    new_entry_id,
    recompute_ledger_workbook,
    transaction_delta,
    utc_now_iso,
)


class ExcelWriter:
    def __init__(self, path: str) -> None:
        self.path = path

    def append_transaction_and_adjust_capital(
        self,
        *,
        txn_date: date,
        description: str,
        category: str,
        txn_type: str,
        whose: str,
        amount: float,
        account: str,
        mandatory: str,
        capital_month: str | None = None,
        capital_delta: float | None = None,
    ) -> tuple[int, float | None]:
        """Append an auditable transaction and rebuild materialized balances."""
        del capital_month, capital_delta
        ensure_ledger_schema(self.path)
        wb = openpyxl.load_workbook(self.path)
        n, _ = self._append_ledger_row(
            wb,
            txn_date=txn_date,
            description=description,
            category=category,
            txn_type=txn_type,
            whose=whose,
            amount=amount,
            account=account,
            mandatory=mandatory,
            kind="transaction",
            delta=transaction_delta(txn_type, amount),
        )
        new_capital = recompute_ledger_workbook(wb)
        wb.save(self.path)
        return n, new_capital

    def _append_ledger_row(
        self,
        wb: Any,
        *,
        txn_date: date,
        description: str,
        category: str,
        txn_type: str,
        whose: str,
        amount: float,
        account: str,
        mandatory: str,
        kind: str,
        delta: float,
        reverses_entry_id: str = "",
    ) -> tuple[int, str]:
        ws = wb[SHEET_TRANSACTIONS]
        new_row = 3
        for row_idx in range(ws.max_row, 2, -1):
            if ws.cell(row_idx, 1).value is not None:
                new_row = row_idx + 1
                break
        entry_id = new_entry_id()
        values = (
            txn_date.strftime("%Y-%m-%d"),
            f"{txn_date.year}-{txn_date.month:02d}",
            description,
            category,
            txn_type,
            whose,
            float(amount),
            account,
            mandatory,
        )
        for col_idx, value in enumerate(values, start=1):
            ws.cell(new_row, col_idx).value = value
        self._copy_derived_formulas(ws, new_row, columns=(10, 11))
        ws.cell(new_row, LEDGER_COL_ENTRY_ID).value = entry_id
        ws.cell(new_row, LEDGER_COL_POSTED_AT).value = utc_now_iso()
        ws.cell(new_row, LEDGER_COL_KIND).value = kind
        ws.cell(new_row, LEDGER_COL_REVERSES_ID).value = reverses_entry_id or None
        ws.cell(new_row, LEDGER_COL_DELTA).value = round(float(delta), 2)
        return new_row, entry_id

    @staticmethod
    def _find_entry_row(ws: Any, entry_id_or_row: str | int) -> int:
        if isinstance(entry_id_or_row, int) or str(entry_id_or_row).isdigit():
            row_num = int(entry_id_or_row)
            if row_num >= 3 and ws.cell(row_num, 1).value:
                return row_num
        needle = str(entry_id_or_row)
        for row_idx in range(3, ws.max_row + 1):
            if str(ws.cell(row_idx, LEDGER_COL_ENTRY_ID).value or "") == needle:
                return row_idx
        raise ValueError(f"Unknown ledger entry: {entry_id_or_row}")

    def _append_reversal(self, wb: Any, entry_id_or_row: str | int) -> tuple[int, str]:
        ws = wb[SHEET_TRANSACTIONS]
        row_num = self._find_entry_row(ws, entry_id_or_row)
        original_id = str(ws.cell(row_num, LEDGER_COL_ENTRY_ID).value or "")
        if not original_id:
            raise ValueError("Ledger entry has no stable ID")
        for existing in range(3, ws.max_row + 1):
            if (
                str(ws.cell(existing, LEDGER_COL_KIND).value or "") == "reversal"
                and str(ws.cell(existing, LEDGER_COL_REVERSES_ID).value or "") == original_id
            ):
                raise ValueError("Ledger entry is already reversed")
        raw_date = ws.cell(row_num, 1).value
        txn_date = raw_date if isinstance(raw_date, date) else date.fromisoformat(str(raw_date)[:10])
        delta = -float(ws.cell(row_num, LEDGER_COL_DELTA).value or 0)
        return self._append_ledger_row(
            wb,
            txn_date=txn_date,
            description=str(ws.cell(row_num, 3).value or ""),
            category=str(ws.cell(row_num, 4).value or ""),
            txn_type=str(ws.cell(row_num, 5).value or ""),
            whose=str(ws.cell(row_num, 6).value or ""),
            amount=float(ws.cell(row_num, 7).value or 0),
            account=str(ws.cell(row_num, 8).value or ""),
            mandatory=str(ws.cell(row_num, 9).value or "Нет"),
            kind="reversal",
            delta=delta,
            reverses_entry_id=original_id,
        )

    def reverse_transaction(self, entry_id_or_row: str | int) -> str:
        ensure_ledger_schema(self.path)
        wb = openpyxl.load_workbook(self.path)
        _, reversal_id = self._append_reversal(wb, entry_id_or_row)
        recompute_ledger_workbook(wb)
        wb.save(self.path)
        return reversal_id

    def correct_transaction(
        self, entry_id_or_row: str | int, txn_date: date, description: str,
        category: str, txn_type: str, whose: str, amount: float,
        account: str, mandatory: str,
    ) -> str:
        ensure_ledger_schema(self.path)
        wb = openpyxl.load_workbook(self.path)
        self._append_reversal(wb, entry_id_or_row)
        _, replacement_id = self._append_ledger_row(
            wb,
            txn_date=txn_date,
            description=description,
            category=category,
            txn_type=txn_type,
            whose=whose,
            amount=amount,
            account=account,
            mandatory=mandatory,
            kind="transaction",
            delta=transaction_delta(txn_type, amount),
        )
        recompute_ledger_workbook(wb)
        wb.save(self.path)
        return replacement_id

    def reconcile_capital(self, txn_date: date, target_balance: float) -> float:
        ensure_ledger_schema(self.path)
        wb = openpyxl.load_workbook(self.path)
        month = txn_date.strftime("%Y-%m")
        ws_capital = wb[SHEET_CAPITAL]
        if not any(
            row[0].value and str(row[0].value).strip() == month
            for row in ws_capital.iter_rows(min_row=3)
        ):
            raise ValueError(f"Month {month} not found in {SHEET_CAPITAL}")
        current = recompute_ledger_workbook(wb, txn_date)
        delta = round(float(target_balance) - current, 2)
        if delta:
            self._append_ledger_row(
                wb,
                txn_date=txn_date,
                description="Сверка капитала",
                category="Корректировка",
                txn_type="Доход" if delta > 0 else "Расход",
                whose="",
                amount=abs(delta),
                account="",
                mandatory="Нет",
                kind="adjustment",
                delta=delta,
            )
        recompute_ledger_workbook(wb)
        wb.save(self.path)
        return float(target_balance)

    def append_transaction(
        self,
        txn_date: date,
        description: str,
        category: str,
        txn_type: str,
        whose: str,
        amount: float,
        account: str,
        mandatory: str,
    ) -> int:
        """Append a transaction row. Returns the new row number."""
        row_num, _ = self.append_transaction_and_adjust_capital(
            txn_date=txn_date,
            description=description,
            category=category,
            txn_type=txn_type,
            whose=whose,
            amount=amount,
            account=account,
            mandatory=mandatory,
        )
        return row_num

    @staticmethod
    def _copy_derived_formulas(ws: Any, target_row: int, columns: tuple[int, ...]) -> None:
        """Copy workbook-template formulas for derived transaction columns."""
        for col_idx in columns:
            formula: str | None = None
            origin_row: int | None = None
            for row_idx in range(target_row - 1, 2, -1):
                value = ws.cell(row_idx, col_idx).value
                if isinstance(value, str) and value.startswith("="):
                    formula = value
                    origin_row = row_idx
                    break

            if formula is None or origin_row is None:
                ws.cell(target_row, col_idx).value = None
                continue

            origin = ws.cell(origin_row, col_idx).coordinate
            target = ws.cell(target_row, col_idx).coordinate
            ws.cell(target_row, col_idx).value = Translator(formula, origin=origin).translate_formula(target)

    def update_capital_salary(
        self,
        months: list[str],
        salary: float,
        ndfl_rate: float = 0.13,
        mandatory_payments: dict[str, list] | None = None,
        pay_days: tuple[int, int] = (5, 20),
        payment_model: str = "working_days",
        payment_percentages: tuple[float, float] = (0.5, 0.5),
    ) -> tuple[list[str], dict[str, float]]:
        # Read current balance for first (current) month so we can project future ones.
        current_balance = 0.0
        if months:
            wb_ro = openpyxl.load_workbook(self.path, data_only=True)
            if SHEET_CAPITAL in wb_ro.sheetnames:
                for row in wb_ro[SHEET_CAPITAL].iter_rows(min_row=3, values_only=True):
                    if row[0] and str(row[0]).strip() == months[0]:
                        try:
                            current_balance = float(row[1] or 0)
                        except (TypeError, ValueError):
                            current_balance = 0.0
                        break

        wb = openpyxl.load_workbook(self.path)
        ws = wb[SHEET_CAPITAL]

        month_rows: dict[str, int] = {}
        for row in ws.iter_rows(min_row=3):
            cell_a = row[0]
            if cell_a.value and isinstance(cell_a.value, str):
                month_rows[cell_a.value.strip()] = cell_a.row

        updated: list[str] = []
        projected: dict[str, float] = {}
        running_balance = current_balance

        for i, month_str in enumerate(months):
            if month_str not in month_rows:
                continue
            row_idx = month_rows[month_str]
            y, m = int(month_str[:4]), int(month_str[5:7])

            try:
                lo, hi = pay_days
                payments = mandatory_payments or {}
                info_lo = calculate_payment(
                    date(y, m, lo),
                    salary,
                    payments.get("first", []),
                    ndfl_rate,
                    pay_days,
                    payment_model,
                    payment_percentages,
                )
                info_hi = calculate_payment(
                    date(y, m, hi),
                    salary,
                    payments.get("second", []),
                    ndfl_rate,
                    pay_days,
                    payment_model,
                    payment_percentages,
                )
                if payment_model == "fixed_percent":
                    coeff_lo = round(payment_percentages[0], 4)
                    coeff_hi = round(payment_percentages[1], 4)
                else:
                    coeff_lo = round(info_lo.working_days_period / info_lo.working_days_month, 4)
                    coeff_hi = round(info_hi.working_days_period / info_hi.working_days_month, 4)
            except ValueError:
                continue

            ws.cell(row_idx, 12).value = round(info_lo.net_amount, 2)
            ws.cell(row_idx, 13).value = round(info_hi.net_amount, 2)
            ws.cell(row_idx, 14).value = coeff_lo
            ws.cell(row_idx, 15).value = coeff_hi

            if i > 0:
                # Accumulate salary remainder (after mandatory payments) for future months.
                running_balance = round(running_balance + info_lo.remainder + info_hi.remainder, 2)
                # Column B is reserved for actual capital. Keep future projections out of it
                # so they do not become "actual" after the calendar month rolls over.
                ws.cell(row_idx, 2).value = None

            projected[month_str] = running_balance
            updated.append(month_str)

        wb.save(self.path)
        return updated, projected

    def update_transaction_row(self, row_num: int, txn_date: date, description: str,
                               category: str, txn_type: str, whose: str,
                               amount: float, account: str, mandatory: str) -> None:
        """Compatibility wrapper: preserve history via reversal + replacement."""
        self.correct_transaction(
            row_num, txn_date, description, category, txn_type, whose,
            amount, account, mandatory,
        )

    def delete_transaction_row(self, row_num: int) -> None:
        """Compatibility wrapper: preserve history by appending a reversal."""
        self.reverse_transaction(row_num)

    def update_setting(self, key: str, value: Any) -> None:
        """Update a key-value row in the Settings sheet. Raises KeyError if key not found."""
        wb = openpyxl.load_workbook(self.path)
        ws = wb[SHEET_SETTINGS]
        for row in ws.iter_rows(min_row=1):
            if row[0].value is not None and str(row[0].value).strip() == key:
                row[1].value = value
                wb.save(self.path)
                return
        raise KeyError(f"Setting '{key}' not found in {SHEET_SETTINGS}")
