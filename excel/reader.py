from __future__ import annotations
import logging
from datetime import date, datetime
from typing import Any
import openpyxl

logger = logging.getLogger(__name__)

from excel.constants import (
    LEDGER_COL_BALANCE_AFTER,
    LEDGER_COL_DELTA,
    LEDGER_COL_ENTRY_ID,
    LEDGER_COL_KIND,
    LEDGER_COL_POSTED_AT,
    LEDGER_COL_REVERSES_ID,
    SETTING_LEDGER_OPENING_BALANCE,
    SETTING_LEDGER_OPENING_DATE,
    SHEET_TRANSACTIONS,
    SHEET_SETTINGS,
    SHEET_CREDITS,
    SHEET_PAYMENTS,
    SHEET_CAPITAL,
)


class ExcelReader:
    def __init__(self, path: str) -> None:
        self.path = path
        self._wb: openpyxl.Workbook | None = None
        self._txns_cache: list[dict[str, Any]] | None = None
        self._txns_by_month_cache: dict[str, list[dict[str, Any]]] | None = None
        self._txns_sorted_cache: list[dict[str, Any]] | None = None
        self._categories_cache: list[str] | None = None
        self._accounts_cache: list[str] | None = None
        self._ledger_entries_cache: list[dict[str, Any]] | None = None

    def _load(self) -> openpyxl.Workbook:
        if self._wb is None:
            self._wb = openpyxl.load_workbook(self.path, data_only=True)
        return self._wb

    def invalidate_cache(self) -> None:
        self._wb = None
        self._txns_cache = None
        self._txns_by_month_cache = None
        self._txns_sorted_cache = None
        self._categories_cache = None
        self._accounts_cache = None
        self._ledger_entries_cache = None

    def _ensure_transactions_cache(self) -> None:
        if self._txns_cache is not None:
            return

        wb = self._load()
        ws = wb[SHEET_TRANSACTIONS]

        all_entries: list[dict[str, Any]] = []
        by_month: dict[str, list[dict[str, Any]]] = {}
        categories_seen: set[str] = set()
        accounts_seen: set[str] = set()
        categories: list[str] = []
        accounts: list[str] = []

        for row_idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            if not row[0]:
                continue
            raw_date = row[0]
            if isinstance(raw_date, datetime):
                txn_date = raw_date.date()
            elif isinstance(raw_date, date):
                txn_date = raw_date
            elif isinstance(raw_date, str):
                try:
                    txn_date = date.fromisoformat(raw_date)
                except ValueError:
                    continue
            else:
                continue

            txn_month = f"{txn_date.year}-{txn_date.month:02d}"
            txn = {
                "date": txn_date,
                "month": txn_month,
                "description": row[2] or "",
                "category": row[3] or "",
                "type": row[4] or "",
                "whose": row[5] or "",
                "amount": float(row[6] or 0),
                "account": row[7] or "",
                "mandatory": row[8] or "Нет",
                "row": row_idx,
                "entry_id": str(row[LEDGER_COL_ENTRY_ID - 1] or "") if len(row) >= LEDGER_COL_ENTRY_ID else "",
                "posted_at": str(row[LEDGER_COL_POSTED_AT - 1] or "") if len(row) >= LEDGER_COL_POSTED_AT else "",
                "entry_kind": str(row[LEDGER_COL_KIND - 1] or "transaction") if len(row) >= LEDGER_COL_KIND else "transaction",
                "reverses_entry_id": str(row[LEDGER_COL_REVERSES_ID - 1] or "") if len(row) >= LEDGER_COL_REVERSES_ID else "",
                "capital_delta": (
                    float(row[LEDGER_COL_DELTA - 1] or 0)
                    if len(row) >= LEDGER_COL_DELTA
                    else (float(row[6] or 0) if row[4] == "Доход" else -float(row[6] or 0))
                ),
                "balance_after": (
                    float(row[LEDGER_COL_BALANCE_AFTER - 1])
                    if len(row) >= LEDGER_COL_BALANCE_AFTER and row[LEDGER_COL_BALANCE_AFTER - 1] is not None
                    else None
                ),
            }
            all_entries.append(txn)

        reversed_ids = {
            entry["reverses_entry_id"]
            for entry in all_entries
            if entry["entry_kind"] == "reversal" and entry["reverses_entry_id"]
        }
        txns = [
            entry for entry in all_entries
            if entry["entry_kind"] == "transaction" and entry["entry_id"] not in reversed_ids
        ]
        for txn in txns:
            by_month.setdefault(txn["month"], []).append(txn)

            cat = txn["category"]
            if cat and cat not in categories_seen:
                categories_seen.add(cat)
                categories.append(cat)

            acc = txn["account"]
            if acc and acc not in accounts_seen:
                accounts_seen.add(acc)
                accounts.append(acc)

        self._txns_cache = txns
        self._txns_by_month_cache = by_month
        self._txns_sorted_cache = sorted(txns, key=lambda t: (t["date"], t["row"]), reverse=True)
        self._categories_cache = categories
        self._accounts_cache = accounts
        self._ledger_entries_cache = all_entries

    def get_transactions(self, month: str | None = None) -> list[dict[str, Any]]:
        self._ensure_transactions_cache()
        if month:
            return list((self._txns_by_month_cache or {}).get(month, []))
        return list(self._txns_cache or [])

    def get_last_transactions(self, n: int = 5) -> list[dict[str, Any]]:
        self._ensure_transactions_cache()
        txns = self._txns_sorted_cache or []
        return list(txns[:n])

    def get_categories(self) -> list[str]:
        self._ensure_transactions_cache()
        return list(self._categories_cache or [])

    def get_accounts(self) -> list[str]:
        self._ensure_transactions_cache()
        return list(self._accounts_cache or [])

    def get_ledger_entries(self) -> list[dict[str, Any]]:
        """Return all raw ledger entries, including reversals and adjustments."""
        self._ensure_transactions_cache()
        return list(self._ledger_entries_cache or [])

    def get_balance_at_start(self, target: date) -> float | None:
        """Return aggregate RUB balance at the start of ``target``.

        The target day's entries are intentionally excluded.  ``None`` means
        the requested date predates the migrated opening anchor.
        """
        settings = self.get_settings()
        opening_raw = settings.get(SETTING_LEDGER_OPENING_DATE)
        balance_raw = settings.get(SETTING_LEDGER_OPENING_BALANCE)
        if opening_raw is None or balance_raw is None:
            return None
        opening_date = date.fromisoformat(str(opening_raw)[:10])
        if target < opening_date:
            return None
        balance = float(balance_raw)
        for entry in self.get_ledger_entries():
            if opening_date <= entry["date"] < target:
                balance += float(entry.get("capital_delta") or 0)
        return round(balance, 2)

    def get_projected_capital_at_start(
        self,
        target: date,
        *,
        reference_date: date,
    ) -> float | None:
        """Project capital at the start of a future month.

        The current month's materialized capital is the anchor. Planned salary
        and mandatory-payment flows are applied only for complete intervening
        months, so a projection for December includes flows through November.
        """
        if target <= reference_date:
            return None
        target_month = target.replace(day=1)
        current_month = reference_date.replace(day=1)
        capital = self.get_capital_data()
        anchor = capital.get(current_month.strftime("%Y-%m"))
        if not anchor:
            return None
        asset_keys = ("rubles", "usd", "investments", "crypto")
        if not any(anchor.get(key) is not None for key in (*asset_keys, "debts")):
            return None
        balance = sum(float(anchor.get(key) or 0) for key in asset_keys)
        balance -= float(anchor.get("debts") or 0)

        payments = self.get_mandatory_payments()
        mandatory_total = sum(
            float(item.get("amount") or 0)
            for half in ("first", "second")
            for item in payments.get(half, [])
        )

        year, month = current_month.year, current_month.month + 1
        if month == 13:
            year, month = year + 1, 1
        cursor = date(year, month, 1)
        while cursor < target_month:
            row = capital.get(cursor.strftime("%Y-%m"))
            if row is None or (
                row.get("salary_first") is None and row.get("salary_second") is None
            ):
                return None
            balance += float(row.get("salary_first") or 0)
            balance += float(row.get("salary_second") or 0)
            balance -= mandatory_total
            year, month = cursor.year, cursor.month + 1
            if month == 13:
                year, month = year + 1, 1
            cursor = date(year, month, 1)
        return round(balance, 2)

    def get_transactions_derived(
        self,
        *,
        current_month: str,
        previous_month: str,
        recent_n: int = 20,
    ) -> dict[str, Any]:
        """Build commonly-needed transaction-derived views from a single scan."""
        self._ensure_transactions_cache()
        current_txns = (self._txns_by_month_cache or {}).get(current_month, [])
        prev_txns = (self._txns_by_month_cache or {}).get(previous_month, [])
        return {
            "current_month_summary": self._summary_from_txns(current_txns, current_month),
            "previous_month_summary": self._summary_from_txns(prev_txns, previous_month),
            "recent_transactions": self.get_last_transactions(recent_n),
            "categories": list(self._categories_cache or []),
            "accounts": list(self._accounts_cache or []),
        }

    @staticmethod
    def _summary_from_txns(txns: list[dict[str, Any]], month: str) -> dict[str, Any]:
        income = sum(t["amount"] for t in txns if t["type"] == "Доход")
        expenses: dict[str, float] = {}
        for t in txns:
            if t["type"] == "Расход":
                expenses[t["category"]] = expenses.get(t["category"], 0) + t["amount"]
        total_expenses = sum(expenses.values())
        return {
            "month": month,
            "income": income,
            "expenses_by_category": expenses,
            "total_expenses": total_expenses,
            "balance": income - total_expenses,
        }

    def get_credits(self) -> list[dict[str, Any]]:
        wb = self._load()
        ws = wb[SHEET_CREDITS]
        result = []
        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row[0] or row[0] == "ИТОГО ДОЛГ":
                continue
            result.append({
                "name": row[0],
                "type": row[1],
                "balance": _numeric_cell(row[2]),
                "monthly_payment": _numeric_cell(row[3]),
                "rate": _numeric_cell(row[4]),
            })
        return result

    def get_settings(self) -> dict[str, Any]:
        """Read key settings: salaries, shares, savings params."""
        wb = self._load()
        ws = wb[SHEET_SETTINGS]
        result: dict[str, Any] = {}
        for row in ws.iter_rows(min_row=1, values_only=True):
            if row[0] and row[1] is not None:
                result[str(row[0]).strip()] = row[1]
        return result

    def get_mandatory_payments(self) -> dict[str, list[dict[str, Any]]]:
        """Return mandatory payments split by half: 'first' (paid ~5th) and 'second' (paid ~20th)."""
        wb = self._load()
        ws = wb[SHEET_PAYMENTS]
        first_half: list[dict[str, Any]] = []
        second_half: list[dict[str, Any]] = []
        current_half: list[dict[str, Any]] = first_half
        header_cols = {"Платёж / назначение", "ИТОГО", "💰", "💡"}
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            name, day, amount = row[0], row[1], row[2]
            if name is None:
                continue
            name_str = str(name)
            if "ВТОРАЯ ЗАРПЛАТА" in name_str or "2️⃣" in name_str:
                current_half = second_half
                continue
            if any(k in name_str for k in header_cols):
                continue
            amount_val = _payment_amount_cell(amount)
            if amount_val is None:
                if day is not None or amount is not None:
                    # Something was entered in День/Сумма but couldn't be parsed
                    # (e.g. a formula cell with no cached value, or stray text) —
                    # a plain section-header row (both cells empty) stays silent.
                    logger.warning(
                        "%s row %d (%r): unparseable amount %r, skipping",
                        SHEET_PAYMENTS, row_idx, name_str, amount,
                    )
                continue
            due_day = _payment_day_cell(day)
            entry = {
                "description": name_str,
                "amount": amount_val,
                "due_day": due_day,
            }
            current_half.append(entry)
        return {"first": first_half, "second": second_half}

    def get_capital_data(self) -> dict[str, dict[str, float | None]]:
        """Return capital rows from 📈 Капитал (row 3+, skipping title+header).

        Keys per month: rubles, usd, investments, crypto, debts, salary_first, salary_second.
        """
        if SHEET_CAPITAL not in self._load().sheetnames:
            return {}
        ws = self._load()[SHEET_CAPITAL]
        result: dict[str, dict[str, float | None]] = {}
        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row[0]:
                continue
            month = str(row[0]).strip()
            if not month[:4].isdigit():
                continue
            def _f(v: Any) -> float | None:
                try:
                    return float(v) if v is not None else None
                except (TypeError, ValueError):
                    return None
            result[month] = {
                "rubles": _f(row[1]),
                "usd": _f(row[2]),
                "investments": _f(row[3]),
                "crypto": _f(row[4]),
                "debts": _f(row[6]),
                "salary_first": _f(row[11]),
                "salary_second": _f(row[12]),
            }
        return result

    def get_summary_data(self, month: str) -> dict[str, Any]:
        txns = self.get_transactions(month=month)
        return self._summary_from_txns(txns, month)


def _payment_amount_cell(value: Any) -> float | None:
    """Tolerantly parse a 📅 Платежи amount cell. Returns None if unparseable.

    Accepts plain numbers, and spreadsheet strings with a currency suffix
    (₽, руб), thousands separators (spaces, NBSP), and comma decimals.
    A formula cell without a cached value (data_only=True) reads as None
    here and is treated as unparseable, not as zero, so it is never
    silently dropped as a real 0 ₽ payment.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        normalized = (
            value.strip()
            .replace(" ", "")
            .replace(" ", "")
            .replace("₽", "")
            .replace("руб.", "")
            .replace("руб", "")
            .replace(",", ".")
        )
        if not normalized:
            return None
        try:
            return float(normalized)
        except ValueError:
            return None
    return None


def _payment_day_cell(value: Any) -> int:
    """Tolerantly parse a 📅 Платежи due-day cell; defaults to 0 (unspecified)."""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)
    if isinstance(value, str):
        normalized = value.strip().replace(" ", "").replace(" ", "")
        if normalized.isdigit():
            return int(normalized)
    return 0


def _numeric_cell(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        normalized = value.strip().replace("\u00a0", "").replace(" ", "")
        if normalized.endswith("%"):
            normalized = normalized[:-1]
        normalized = normalized.replace(",", ".")
        return float(normalized or 0)
    return float(value or 0)
