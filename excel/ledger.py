"""Append-only Excel ledger helpers.

The transaction journal is the source of truth.  ``balance_after`` and the
current-month capital cell are materialized views and may always be rebuilt
from the opening anchor plus ledger deltas.
"""
from __future__ import annotations

from dataclasses import dataclass
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

import openpyxl

from excel.constants import (
    LEDGER_COL_BALANCE_AFTER,
    LEDGER_COL_DELTA,
    LEDGER_COL_ENTRY_ID,
    LEDGER_COL_KIND,
    LEDGER_COL_POSTED_AT,
    LEDGER_COL_REVERSES_ID,
    LEDGER_HEADERS,
    LEDGER_SCHEMA_VERSION,
    SETTING_LEDGER_OPENING_BALANCE,
    SETTING_LEDGER_OPENING_DATE,
    SETTING_LEDGER_VERSION,
    SHEET_CAPITAL,
    SHEET_SETTINGS,
    SHEET_TRANSACTIONS,
)


@dataclass(frozen=True)
class LedgerMigrationResult:
    changed: bool
    migrated_entries: int = 0
    opening_date: date | None = None
    opening_balance: float | None = None


def new_entry_id() -> str:
    return uuid4().hex


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds")


def transaction_delta(txn_type: str, amount: float) -> float:
    return float(amount) if txn_type == "Доход" else -float(amount)


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return None
    return None


def _settings(ws: Any) -> dict[str, Any]:
    return {
        str(row[0]).strip(): row[1]
        for row in ws.iter_rows(min_row=1, values_only=True)
        if row[0] is not None
    }


def _upsert_setting(ws: Any, key: str, value: Any) -> None:
    for row in ws.iter_rows(min_row=1):
        if row[0].value is not None and str(row[0].value).strip() == key:
            row[1].value = value
            return
    target = ws.max_row + 1
    ws.cell(target, 1).value = key
    ws.cell(target, 2).value = value


def _capital_balance(ws: Any, on_date: date) -> float:
    month = on_date.strftime("%Y-%m")
    previous: tuple[str, float] | None = None
    current: float | None = None
    found = False
    for row in ws.iter_rows(min_row=3, values_only=True):
        label = str(row[0]).strip() if row[0] else ""
        try:
            value = float(row[1] or 0)
        except (TypeError, ValueError):
            value = 0.0
        if label == month:
            current = value
            found = True
        elif label and label < month and (previous is None or label > previous[0]):
            previous = (label, value)
    if found:
        return current
    return previous[1] if previous else 0.0


def _set_current_capital(ws: Any, on_date: date, value: float) -> bool:
    """Write the current-month capital cell. Returns True if the month row was found."""
    month = on_date.strftime("%Y-%m")
    for row in ws.iter_rows(min_row=3):
        if row[0].value and str(row[0].value).strip() == month:
            row[1].value = round(value, 2)
            return True
    return False


def _entry_rows(ws: Any) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        effective_date = _parse_date(row[0] if row else None)
        entry_id = str(row[LEDGER_COL_ENTRY_ID - 1] or "").strip() if len(row) >= LEDGER_COL_ENTRY_ID else ""
        if effective_date is None or not entry_id:
            continue
        try:
            delta = float(row[LEDGER_COL_DELTA - 1] or 0)
        except (TypeError, ValueError):
            delta = 0.0
        entries.append({
            "row": row_idx,
            "date": effective_date,
            "entry_id": entry_id,
            "posted_at": str(row[LEDGER_COL_POSTED_AT - 1] or "") if len(row) >= LEDGER_COL_POSTED_AT else "",
            "kind": str(row[LEDGER_COL_KIND - 1] or "transaction") if len(row) >= LEDGER_COL_KIND else "transaction",
            "reverses_id": str(row[LEDGER_COL_REVERSES_ID - 1] or "") if len(row) >= LEDGER_COL_REVERSES_ID else "",
            "delta": delta,
            "balance_after": (
                float(row[LEDGER_COL_BALANCE_AFTER - 1])
                if len(row) >= LEDGER_COL_BALANCE_AFTER and row[LEDGER_COL_BALANCE_AFTER - 1] is not None
                else None
            ),
        })
    return entries


def recompute_ledger_workbook(wb: Any, on_date: date | None = None) -> float:
    """Rebuild all derived balances and current capital; return today's balance."""
    on_date = on_date or date.today()
    ws_tx = wb[SHEET_TRANSACTIONS]
    settings = _settings(wb[SHEET_SETTINGS])
    opening_date = date.fromisoformat(str(settings[SETTING_LEDGER_OPENING_DATE])[:10])
    opening_balance = float(settings[SETTING_LEDGER_OPENING_BALANCE])
    entries = sorted(
        _entry_rows(ws_tx),
        key=lambda e: (e["date"], e["posted_at"], e["row"]),
    )
    if entries and entries[0]["date"] < opening_date:
        # A newly posted backdated entry extends the best-effort history.  Keep
        # the same opening amount and move only the anchor date backwards so
        # the new delta affects all subsequent balances.
        opening_date = entries[0]["date"]
        _upsert_setting(wb[SHEET_SETTINGS], SETTING_LEDGER_OPENING_DATE, opening_date.isoformat())
    running = opening_balance
    for entry in entries:
        if entry["date"] < opening_date:
            ws_tx.cell(entry["row"], LEDGER_COL_BALANCE_AFTER).value = None
            continue
        running = round(running + entry["delta"], 2)
        ws_tx.cell(entry["row"], LEDGER_COL_BALANCE_AFTER).value = running

    current = round(
        opening_balance
        + sum(e["delta"] for e in entries if opening_date <= e["date"] <= on_date),
        2,
    )
    _set_current_capital(wb[SHEET_CAPITAL], on_date, current)
    return current


def ensure_ledger_schema(path: str | Path, on_date: date | None = None) -> LedgerMigrationResult:
    """Idempotently migrate a runtime workbook to ledger schema v1."""
    on_date = on_date or date.today()
    wb = openpyxl.load_workbook(path)
    ws_tx = wb[SHEET_TRANSACTIONS]
    ws_settings = wb[SHEET_SETTINGS]
    current = _settings(ws_settings)
    if str(current.get(SETTING_LEDGER_VERSION, "")) == LEDGER_SCHEMA_VERSION:
        return LedgerMigrationResult(False)

    for offset, header in enumerate(LEDGER_HEADERS, start=LEDGER_COL_ENTRY_ID):
        ws_tx.cell(2, offset).value = header

    legacy_rows: list[tuple[int, date, float]] = []
    migration_time = datetime.now(timezone.utc)
    for sequence, row_idx in enumerate(range(3, ws_tx.max_row + 1)):
        txn_date = _parse_date(ws_tx.cell(row_idx, 1).value)
        if txn_date is None:
            continue
        amount = float(ws_tx.cell(row_idx, 7).value or 0)
        delta = transaction_delta(str(ws_tx.cell(row_idx, 5).value or ""), amount)
        ws_tx.cell(row_idx, LEDGER_COL_ENTRY_ID).value = new_entry_id()
        posted_at = migration_time + timedelta(microseconds=sequence)
        ws_tx.cell(row_idx, LEDGER_COL_POSTED_AT).value = posted_at.isoformat(timespec="microseconds")
        ws_tx.cell(row_idx, LEDGER_COL_KIND).value = "transaction"
        ws_tx.cell(row_idx, LEDGER_COL_REVERSES_ID).value = None
        ws_tx.cell(row_idx, LEDGER_COL_DELTA).value = delta
        legacy_rows.append((row_idx, txn_date, delta))

    opening_date = min((item[1] for item in legacy_rows), default=on_date)
    capital_now = _capital_balance(wb[SHEET_CAPITAL], on_date)
    included_delta = sum(delta for _, txn_date, delta in legacy_rows if txn_date <= on_date)
    opening_balance = round(capital_now - included_delta, 2)
    _upsert_setting(ws_settings, SETTING_LEDGER_VERSION, LEDGER_SCHEMA_VERSION)
    _upsert_setting(ws_settings, SETTING_LEDGER_OPENING_DATE, opening_date.isoformat())
    _upsert_setting(ws_settings, SETTING_LEDGER_OPENING_BALANCE, opening_balance)
    recompute_ledger_workbook(wb, on_date)
    wb.save(path)
    return LedgerMigrationResult(True, len(legacy_rows), opening_date, opening_balance)


def needs_ledger_migration(path: str | Path) -> bool:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    return str(_settings(wb[SHEET_SETTINGS]).get(SETTING_LEDGER_VERSION, "")) != LEDGER_SCHEMA_VERSION


def validate_ledger_workbook(path: str | Path) -> list[str]:
    """Return structural integrity errors without mutating the workbook."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    settings = _settings(wb[SHEET_SETTINGS])
    errors: list[str] = []
    for key in (SETTING_LEDGER_VERSION, SETTING_LEDGER_OPENING_DATE, SETTING_LEDGER_OPENING_BALANCE):
        if key not in settings:
            errors.append(f"missing setting: {key}")
    entries = _entry_rows(wb[SHEET_TRANSACTIONS])
    ids = [e["entry_id"] for e in entries]
    duplicate_ids = sorted(entry_id for entry_id, count in Counter(ids).items() if count > 1)
    if duplicate_ids:
        errors.append("duplicate entry_id: " + ", ".join(duplicate_ids[:5]))
    known = set(ids)
    for entry in entries:
        if entry["kind"] == "reversal" and entry["reverses_id"] not in known:
            errors.append(f"broken reversal: {entry['entry_id']}")
    try:
        opening_date = date.fromisoformat(str(settings[SETTING_LEDGER_OPENING_DATE])[:10])
        running = float(settings[SETTING_LEDGER_OPENING_BALANCE])
        for entry in sorted(entries, key=lambda item: (item["date"], item["posted_at"], item["row"])):
            if entry["date"] < opening_date:
                continue
            running = round(running + entry["delta"], 2)
            stored = entry["balance_after"]
            if stored is None or abs(float(stored) - running) > 0.01:
                errors.append(f"balance_after mismatch: {entry['entry_id']}")
                break
    except (KeyError, TypeError, ValueError):
        pass
    return errors
