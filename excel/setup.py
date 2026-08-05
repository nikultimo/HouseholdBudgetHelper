from __future__ import annotations

import os
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell

from excel.constants import (
    SHEET_CAPITAL,
    SHEET_CREDITS,
    SHEET_PAYMENTS,
    SHEET_SETTINGS,
    SHEET_TRANSACTIONS,
)


REQUIRED_TEMPLATE_SHEETS = {
    SHEET_TRANSACTIONS,
    SHEET_SETTINGS,
    SHEET_CREDITS,
    SHEET_PAYMENTS,
    SHEET_CAPITAL,
}

_EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
_PHONE_RE = re.compile(r"(?:\+?\d[\s().-]*){10,}")
_TOKEN_RE = re.compile(r"(?:token|api[_-]?key|password|secret)\s*[:=]\s*[\w.-]{8,}", re.I)


@dataclass
class MandatoryPaymentSetup:
    description: str
    amount: float
    due_day: int = 0
    owner: str = "Общее"


@dataclass
class CreditSetup:
    name: str
    balance: float
    monthly_payment: float = 0.0
    credit_type: str = "Кредит"
    rate: float = 0.0


@dataclass
class WorkbookSetupData:
    user_name: str
    default_account: str = "Main Card"
    initial_capital: float = 0.0
    salary: float = 0.0
    extra_income: float = 0.0
    salary_pay_days: tuple[int, int] = (5, 20)
    salary_payment_model: str = "working_days"
    salary_payment_percentages: tuple[float, float] = (0.5, 0.5)
    categories: list[str] = field(default_factory=list)
    accounts: list[str] = field(default_factory=list)
    mandatory_payments: list[MandatoryPaymentSetup] = field(default_factory=list)
    credits: list[CreditSetup] = field(default_factory=list)


def validate_template_workbook(path: str | os.PathLike[str]) -> None:
    """Validate the public workbook template before replacing tracked fixtures."""
    wb = openpyxl.load_workbook(path, data_only=False)
    missing = REQUIRED_TEMPLATE_SHEETS.difference(wb.sheetnames)
    if missing:
        raise ValueError(f"Template is missing required sheets: {', '.join(sorted(missing))}")

    ws_tx = wb[SHEET_TRANSACTIONS]
    if ws_tx.max_row < 2 or ws_tx.max_column < 9:
        raise ValueError(f"{SHEET_TRANSACTIONS} must have headers in rows 1-2 and columns A-I")

    ws_pay = wb[SHEET_PAYMENTS]
    if ws_pay.max_row < 2:
        raise ValueError(f"{SHEET_PAYMENTS} must contain payment rows from row 2")
    if not any(
        cell.value and "ВТОРАЯ ЗАРПЛАТА" in str(cell.value)
        for row in ws_pay.iter_rows(min_row=2)
        for cell in row
    ):
        raise ValueError(f"{SHEET_PAYMENTS} must contain a 'ВТОРАЯ ЗАРПЛАТА' separator")

    ws_cap = wb[SHEET_CAPITAL]
    if ws_cap.max_row < 3:
        raise ValueError(f"{SHEET_CAPITAL} must contain month rows from row 3")
    if not str(ws_cap.cell(3, 1).value or "")[:4].isdigit():
        raise ValueError(f"{SHEET_CAPITAL} row 3 must start capital month rows")
    for col in range(12, 16):
        if ws_cap.cell(2, col).value is None:
            raise ValueError(f"{SHEET_CAPITAL} must define salary projection columns L-O")

    suspicious = scan_workbook_for_private_data(path)
    if suspicious:
        sample = "; ".join(suspicious[:5])
        raise ValueError(f"Template contains suspicious private values: {sample}")


def scan_workbook_for_private_data(path: str | os.PathLike[str]) -> list[str]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    findings: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str):
                    continue
                text = value.strip()
                if not text:
                    continue
                if _EMAIL_RE.search(text) or _PHONE_RE.search(text) or _TOKEN_RE.search(text):
                    findings.append(f"{ws.title}!{cell.coordinate}={text[:80]}")
    return findings


def ensure_runtime_workbook(
    runtime_path: str | os.PathLike[str],
    template_path: str | os.PathLike[str] = "example/budget_example.xlsx",
) -> bool:
    """Create the runtime workbook from the public template if it does not exist."""
    runtime = Path(runtime_path)
    if runtime.exists():
        return False
    validate_template_workbook(template_path)
    runtime.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, runtime)
    return True


def apply_workbook_setup(
    runtime_path: str | os.PathLike[str],
    setup: WorkbookSetupData,
    template_path: str | os.PathLike[str] = "example/budget_example.xlsx",
) -> None:
    ensure_runtime_workbook(runtime_path, template_path)
    wb = openpyxl.load_workbook(runtime_path)

    ws_settings = wb[SHEET_SETTINGS]
    _upsert_setting(ws_settings, "Имя 1", setup.user_name)
    _upsert_setting(ws_settings, "Счёт по умолчанию", setup.default_account)
    _upsert_setting(ws_settings, f"Зарплата {setup.user_name}", float(setup.salary))
    _upsert_setting(ws_settings, f"Доп. доход {setup.user_name}", float(setup.extra_income))
    _upsert_setting(ws_settings, f"Доля {setup.user_name} %", 1.0)
    _upsert_setting(ws_settings, "salary_pay_days", ",".join(str(d) for d in setup.salary_pay_days))
    _upsert_setting(ws_settings, "salary_payment_model", setup.salary_payment_model)
    _upsert_setting(
        ws_settings,
        "salary_payment_percentages",
        ",".join(str(v) for v in setup.salary_payment_percentages),
    )

    categories = [c.strip() for c in setup.categories if c.strip()]
    if categories:
        _replace_list_section(ws_settings, "🏷️", categories)

    accounts = [a.strip() for a in setup.accounts if a.strip()]
    if accounts:
        _upsert_setting(ws_settings, "Счета", ", ".join(accounts))

    if setup.mandatory_payments:
        _write_mandatory_payments(wb[SHEET_PAYMENTS], setup.mandatory_payments, setup.salary_pay_days)

    if setup.credits:
        _write_credits(wb[SHEET_CREDITS], setup.credits)

    if setup.initial_capital:
        _set_first_capital_balance(wb[SHEET_CAPITAL], setup.initial_capital)

    wb.save(runtime_path)


def _upsert_setting(ws: Any, key: str, value: Any) -> None:
    for row in ws.iter_rows(min_row=1):
        if row[0].value is not None and str(row[0].value).strip() == key:
            row[1].value = value
            return
    row_idx = ws.max_row + 1
    ws.cell(row_idx, 1).value = key
    ws.cell(row_idx, 2).value = value


def _replace_list_section(ws: Any, marker: str, values: list[str]) -> None:
    start_row: int | None = None
    for row in ws.iter_rows(min_row=1):
        if row[0].value and marker in str(row[0].value):
            start_row = row[0].row + 1
            break
    if start_row is None:
        start_row = ws.max_row + 1
        ws.cell(start_row - 1, 1).value = "🏷️ КАТЕГОРИИ РАСХОДОВ"

    end_row = start_row
    while end_row <= ws.max_row and ws.cell(end_row, 1).value:
        end_row += 1
    existing_slots = max(0, end_row - start_row)
    for offset in range(max(existing_slots, len(values))):
        cell = ws.cell(start_row + offset, 1)
        cell.value = values[offset] if offset < len(values) else None


def _write_mandatory_payments(
    ws: Any,
    payments: list[MandatoryPaymentSetup],
    pay_days: tuple[int, int],
) -> None:
    first = [p for p in payments if not p.due_day or p.due_day <= pay_days[0]]
    second = [p for p in payments if p.due_day > pay_days[0]]
    for row_idx in range(4, ws.max_row + 1):
        for col_idx in range(1, 6):
            cell = ws.cell(row_idx, col_idx)
            if not isinstance(cell, MergedCell):
                cell.value = None

    row_idx = 4
    for payment in first:
        _write_payment_row(ws, row_idx, payment)
        row_idx += 1

    separator_row = max(row_idx + 1, 16)
    ws.cell(separator_row, 1).value = "2️⃣  ВТОРАЯ ЗАРПЛАТА"
    row_idx = separator_row + 2
    for payment in second:
        _write_payment_row(ws, row_idx, payment)
        row_idx += 1


def _write_payment_row(ws: Any, row_idx: int, payment: MandatoryPaymentSetup) -> None:
    ws.cell(row_idx, 1).value = payment.description
    ws.cell(row_idx, 2).value = payment.due_day
    ws.cell(row_idx, 3).value = float(payment.amount)
    ws.cell(row_idx, 4).value = payment.owner
    ws.cell(row_idx, 5).value = "Да"


def _write_credits(ws: Any, credits: list[CreditSetup]) -> None:
    for row_idx in range(3, ws.max_row + 1):
        for col_idx in range(1, 6):
            ws.cell(row_idx, col_idx).value = None
    for offset, credit in enumerate(credits):
        row_idx = 3 + offset
        ws.cell(row_idx, 1).value = credit.name
        ws.cell(row_idx, 2).value = credit.credit_type
        ws.cell(row_idx, 3).value = float(credit.balance)
        ws.cell(row_idx, 4).value = float(credit.monthly_payment)
        ws.cell(row_idx, 5).value = float(credit.rate)


def _set_first_capital_balance(ws: Any, value: float) -> None:
    for row in ws.iter_rows(min_row=3):
        if row[0].value and str(row[0].value).strip()[:4].isdigit():
            row[1].value = float(value)
            return
