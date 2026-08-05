from datetime import date

import openpyxl
import pytest

from excel.ledger import _capital_balance, ensure_ledger_schema, validate_ledger_workbook
from excel.reader import ExcelReader
from excel.writer import ExcelWriter


def _ledger_workbook(tmp_path):
    path = tmp_path / "ledger.xlsx"
    wb = openpyxl.Workbook()
    ws_tx = wb.active
    ws_tx.title = "📋 Транзакции"
    ws_tx.append([""] * 11)
    ws_tx.append(["Дата", "Месяц", "Описание", "Категория", "Тип", "Кто", "Сумма", "Счёт", "Обязательный", "", ""])
    ws_tx.append(["2026-06-14", "2026-06", "Кофе", "Еда", "Расход", "User1", 100, "Main", "Нет"])
    ws_tx.append(["2026-06-15", "2026-06", "Возврат", "Доход", "Доход", "User1", 50, "Main", "Нет"])
    wb.create_sheet("⚙️ Настройки")
    wb.create_sheet("💳 Кредиты")
    wb.create_sheet("📅 Платежи")
    ws_cap = wb.create_sheet("📈 Капитал")
    ws_cap["A2"] = "Месяц"
    ws_cap["B2"] = "Счета руб"
    ws_cap["A3"] = "2026-07"
    ws_cap["B3"] = 950
    wb.save(path)
    return path


def test_migration_builds_anchor_and_start_of_day_balance(tmp_path):
    path = _ledger_workbook(tmp_path)
    result = ensure_ledger_schema(path, on_date=date(2026, 7, 13))

    assert result.changed is True
    assert result.opening_balance == 1000
    assert ensure_ledger_schema(path, on_date=date(2026, 7, 13)).changed is False
    assert validate_ledger_workbook(path) == []

    reader = ExcelReader(str(path))
    assert reader.get_balance_at_start(date(2026, 6, 14)) == 1000
    assert reader.get_balance_at_start(date(2026, 6, 15)) == 900
    assert reader.get_balance_at_start(date(2026, 6, 16)) == 950


def test_correction_and_delete_are_append_only(tmp_path):
    path = _ledger_workbook(tmp_path)
    ensure_ledger_schema(path, on_date=date(2026, 7, 13))
    reader = ExcelReader(str(path))
    original = reader.get_transactions()[0]
    writer = ExcelWriter(str(path))

    writer.correct_transaction(
        original["entry_id"], date(2026, 6, 14), "Кофе", "Еда",
        "Расход", "User1", 200, "Main", "Нет",
    )
    reader.invalidate_cache()
    active = reader.get_transactions()
    assert [txn["amount"] for txn in active if txn["description"] == "Кофе"] == [200]
    assert [entry["entry_kind"] for entry in reader.get_ledger_entries()].count("reversal") == 1
    assert reader.get_balance_at_start(date(2026, 6, 15)) == 800

    replacement = next(txn for txn in active if txn["description"] == "Кофе")
    writer.reverse_transaction(replacement["entry_id"])
    reader.invalidate_cache()
    assert all(txn["description"] != "Кофе" for txn in reader.get_transactions())
    assert reader.get_balance_at_start(date(2026, 6, 15)) == 1000


def test_backdated_entry_moves_anchor_and_affects_current_balance(tmp_path):
    path = _ledger_workbook(tmp_path)
    ensure_ledger_schema(path, on_date=date(2026, 7, 13))
    writer = ExcelWriter(str(path))
    writer.append_transaction_and_adjust_capital(
        txn_date=date(2026, 5, 1), description="Ранняя трата", category="Прочее",
        txn_type="Расход", whose="User1", amount=25, account="Main",
        mandatory="Нет",
    )
    reader = ExcelReader(str(path))
    assert reader.get_balance_at_start(date(2026, 5, 2)) == 975


def test_reconciliation_is_auditable_adjustment(tmp_path):
    path = _ledger_workbook(tmp_path)
    ensure_ledger_schema(path, on_date=date(2026, 7, 13))

    ExcelWriter(str(path)).reconcile_capital(date(2026, 7, 13), 1200)

    reader = ExcelReader(str(path))
    adjustments = [entry for entry in reader.get_ledger_entries() if entry["entry_kind"] == "adjustment"]
    assert len(adjustments) == 1
    assert adjustments[0]["capital_delta"] == 250
    assert reader.get_balance_at_start(date(2026, 7, 14)) == 1200


def test_reconcile_capital_raises_when_month_missing(tmp_path):
    path = _ledger_workbook(tmp_path)
    ensure_ledger_schema(path, on_date=date(2026, 7, 13))
    writer = ExcelWriter(str(path))

    with pytest.raises(ValueError):
        writer.reconcile_capital(date(2026, 9, 1), 1500)


def test_reverse_transaction_twice_raises(tmp_path):
    path = _ledger_workbook(tmp_path)
    ensure_ledger_schema(path, on_date=date(2026, 7, 13))
    reader = ExcelReader(str(path))
    entry_id = reader.get_transactions()[0]["entry_id"]
    writer = ExcelWriter(str(path))

    writer.reverse_transaction(entry_id)
    with pytest.raises(ValueError):
        writer.reverse_transaction(entry_id)


def test_capital_balance_treats_zero_as_present(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "📈 Капитал"
    ws.append([""] * 2)
    ws.append(["Месяц", "Счета руб"])
    ws.append(["2026-07", 950])
    ws.append(["2026-08", 0])

    assert _capital_balance(ws, date(2026, 8, 15)) == 0.0
    assert _capital_balance(ws, date(2026, 7, 15)) == 950.0


def test_integrity_check_detects_derived_balance_corruption(tmp_path):
    path = _ledger_workbook(tmp_path)
    ensure_ledger_schema(path, on_date=date(2026, 7, 13))
    wb = openpyxl.load_workbook(path)
    wb["📋 Транзакции"].cell(3, 17).value = -999
    wb.save(path)

    assert any("balance_after mismatch" in error for error in validate_ledger_workbook(path))
