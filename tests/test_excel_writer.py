import shutil
import pytest
import openpyxl
import tempfile
import os
from datetime import date
from pathlib import Path
from unittest.mock import patch
from excel.reader import ExcelReader
from excel.writer import ExcelWriter
from salary.calculator import PaymentInfo

SAMPLE_FILE = Path("example/budget_example.xlsx")
SETTINGS_SHEET = "⚙️ Настройки"


def _make_workbook_with_settings(settings: dict) -> str:
    wb = openpyxl.Workbook()
    ws_t = wb.active
    ws_t.title = "📋 Транзакции"
    ws_t.append([""] * 11)
    ws_t.append(["Дата"] + [""] * 10)
    ws_s = wb.create_sheet(SETTINGS_SHEET)
    for k, v in settings.items():
        ws_s.append([k, v])
    path = tempfile.mktemp(suffix=".xlsx")
    wb.save(path)
    return path


def _make_workbook_with_capital(month: str, balance: float) -> str:
    wb = openpyxl.Workbook()
    ws_t = wb.active
    ws_t.title = "📋 Транзакции"
    ws_t.append([""] * 11)
    ws_t.append(["Дата"] + [""] * 10)

    ws_cap = wb.create_sheet("📈 Капитал")
    ws_cap["A2"] = "Месяц"
    ws_cap["B2"] = "Счета руб"
    ws_cap.cell(3, 1).value = month
    ws_cap.cell(3, 2).value = balance

    # Writer formulas reference settings sheet; keep it present.
    ws_set = wb.create_sheet("⚙️ Настройки")
    ws_set["A15"] = "Доля User1"
    ws_set["B15"] = 0.5
    ws_set["A16"] = "Доля User2"
    ws_set["B16"] = 0.5

    path = tempfile.mktemp(suffix=".xlsx")
    wb.save(path)
    return path


@pytest.fixture
def budget_path(tmp_path):
    dest = tmp_path / "budget.xlsx"
    shutil.copy(SAMPLE_FILE, dest)
    return str(dest)


def test_writer_appends_transaction(budget_path):
    reader = ExcelReader(budget_path)
    count_before = len(reader.get_transactions())

    writer = ExcelWriter(budget_path)
    writer.append_transaction(
        txn_date=date(2026, 4, 16),
        description="Тест кофе",
        category="Еда и продукты",
        txn_type="Расход",
        whose="User1",
        amount=350.0,
        account="Main Card",
        mandatory="Нет",
    )

    reader2 = ExcelReader(budget_path)
    txns_after = reader2.get_transactions()
    assert len(txns_after) == count_before + 1
    last = txns_after[-1]
    assert last["description"] == "Тест кофе"
    assert last["amount"] == 350.0
    assert last["category"] == "Еда и продукты"


def test_writer_sets_correct_formulas(budget_path):
    import openpyxl
    writer = ExcelWriter(budget_path)
    # Use the returned row number — ws.max_row is unreliable because the
    # template has pre-seeded formula rows beyond the last data row.
    last_row = writer.append_transaction(
        txn_date=date(2026, 4, 16),
        description="Формула тест",
        category="Бензин",
        txn_type="Расход",
        whose="User1",
        amount=1000.0,
        account="Main Card",
        mandatory="Да",
    )
    wb = openpyxl.load_workbook(budget_path)
    ws = wb["📋 Транзакции"]
    # Column B should be a static YYYY-MM string (not a formula)
    assert ws.cell(last_row, 2).value == "2026-04"
    # Column J should be a formula
    assert str(ws.cell(last_row, 10).value).startswith(f'=IF(G{last_row}')
    assert "User1" in str(ws.cell(last_row, 10).value)
    # Column K should be a formula
    assert str(ws.cell(last_row, 11).value).startswith(f'=IF(G{last_row}')
    assert "User2" in str(ws.cell(last_row, 11).value)


def test_writer_returns_row_number(budget_path):
    writer = ExcelWriter(budget_path)
    row_num = writer.append_transaction(
        txn_date=date(2026, 4, 16),
        description="Row number test",
        category="Всякое",
        txn_type="Расход",
        whose="User1",
        amount=100.0,
        account="Main Card",
        mandatory="Нет",
    )
    assert isinstance(row_num, int)
    assert row_num >= 3


def test_writer_deletes_transaction_row():
    path = _make_workbook_with_capital(month="2026-04", balance=1000.0)
    try:
        writer = ExcelWriter(path)
        row_num = writer.append_transaction(
            txn_date=date(2026, 4, 16),
            description="Delete row test",
            category="Всякое",
            txn_type="Расход",
            whose="User1",
            amount=100.0,
            account="Main Card",
            mandatory="Нет",
        )

        writer.delete_transaction_row(row_num)

        reader = ExcelReader(path)
        assert all(t["description"] != "Delete row test" for t in reader.get_transactions())
    finally:
        os.unlink(path)


def test_append_transaction_and_adjust_capital_updates_capital_cell():
    path = _make_workbook_with_capital(month="2026-04", balance=1000.0)
    try:
        writer = ExcelWriter(path)
        row_num, new_capital = writer.append_transaction_and_adjust_capital(
            txn_date=date(2026, 4, 16),
            description="Кофе",
            category="Еда",
            txn_type="Расход",
            whose="User1",
            amount=200.0,
            account="Main Card",
            mandatory="Нет",
            capital_month="2026-04",
            capital_delta=-200.0,
        )
        assert row_num >= 3
        assert new_capital == 800.0

        wb = openpyxl.load_workbook(path, data_only=True)
        ws_tx = wb["📋 Транзакции"]
        assert ws_tx.cell(row_num, 17).value == 800.0
    finally:
        os.unlink(path)


def _make_payment_info(net: float, period_days: int, month_days: int, mandatory_total: float = 0.0) -> PaymentInfo:
    return PaymentInfo(
        payment_date=date(2026, 5, 5),
        period_start=date(2026, 4, 16),
        period_end=date(2026, 4, 30),
        working_days_period=period_days,
        working_days_month=month_days,
        monthly_salary=100000.0,
        gross_amount=round(net / 0.87, 2),
        net_amount=net,
        mandatory_payments=[],
        mandatory_total=mandatory_total,
        remainder=round(net - mandatory_total, 2),
    )


@pytest.fixture
def capital_budget_path(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws_set = wb.create_sheet("⚙️ Настройки")
    ws_set["A1"] = "Зарплата User1"
    ws_set["B1"] = 100000.0
    wb.create_sheet("📋 Транзакции")
    ws_cap = wb.create_sheet("📈 Капитал")
    ws_cap["A2"] = "Месяц"
    ws_cap["L2"] = "ЗП 5 ого числа"
    ws_cap["M2"] = "ЗП 20 числа"
    ws_cap["N2"] = "Коэфф зп"
    ws_cap["O2"] = "Коэфф аванса (20 число)"
    for i, month in enumerate(["2026-04", "2026-05", "2026-06"]):
        ws_cap.cell(3 + i, 1).value = month
    ws_cap.cell(4, 2).value = 123456.0
    path = tmp_path / "capital_test.xlsx"
    wb.save(str(path))
    return str(path)


def test_update_capital_salary_writes_values(capital_budget_path):
    # mandatory_total=5000 for 5th payment, 3000 for 20th → remainder used for projection
    info5 = _make_payment_info(net=47619.0, period_days=10, month_days=21, mandatory_total=5000.0)
    info20 = _make_payment_info(net=47619.0, period_days=10, month_days=21, mandatory_total=3000.0)
    mandatory_payments = {
        "first": [{"description": "Test", "amount": 5000.0, "due_day": 5}],
        "second": [{"description": "Test2", "amount": 3000.0, "due_day": 20}],
    }

    with patch("excel.writer.calculate_payment", side_effect=[info5, info20, info5, info20]):
        writer = ExcelWriter(capital_budget_path)
        updated, projected = writer.update_capital_salary(
            ["2026-04", "2026-05"], 100000.0, mandatory_payments=mandatory_payments
        )

    assert updated == ["2026-04", "2026-05"]
    wb = openpyxl.load_workbook(capital_budget_path)
    ws = wb["📈 Капитал"]
    # Row 3 = 2026-04: salary cols written, col B not overwritten (current month)
    assert ws.cell(3, 12).value == 47619.0   # L
    assert ws.cell(3, 13).value == 47619.0   # M
    assert ws.cell(3, 14).value == round(10 / 21, 4)  # N
    assert ws.cell(3, 15).value == round(10 / 21, 4)  # O
    assert ws.cell(3, 2).value is None
    # Row 4 = 2026-05: projection is returned but not written to actual-capital col B.
    expected_projection = round((47619.0 - 5000.0) + (47619.0 - 3000.0), 2)
    assert projected["2026-05"] == expected_projection
    assert ws.cell(4, 2).value is None


def test_update_capital_salary_skips_missing_month(capital_budget_path):
    info5 = _make_payment_info(net=47619.0, period_days=10, month_days=21)
    info20 = _make_payment_info(net=47619.0, period_days=10, month_days=21)

    with patch("excel.writer.calculate_payment", side_effect=[info5, info20]):
        writer = ExcelWriter(capital_budget_path)
        updated, projected = writer.update_capital_salary(["2026-04", "2099-01"], 100000.0)

    assert updated == ["2026-04"]
    assert "2099-01" not in updated


def test_update_capital_salary_skips_month_on_calendar_error(capital_budget_path):
    info_ok = _make_payment_info(net=47619.0, period_days=10, month_days=21)

    def side_effect(*args, **kwargs):
        d = args[0]
        if d.month == 5:
            raise ValueError("No working days found in month")
        return info_ok

    with patch("excel.writer.calculate_payment", side_effect=side_effect):
        writer = ExcelWriter(capital_budget_path)
        updated, projected = writer.update_capital_salary(["2026-04", "2026-05", "2026-06"], 100000.0)

    # 2026-05 skipped due to ValueError, others updated
    assert "2026-04" in updated
    assert "2026-05" not in updated
    assert "2026-06" in updated


def test_update_setting_changes_value():
    path = _make_workbook_with_settings({"Зарплата User1": 100000, "Капитал": 500000})
    try:
        writer = ExcelWriter(path)
        writer.update_setting("Зарплата User1", 120000)
        wb = openpyxl.load_workbook(path)
        ws = wb[SETTINGS_SHEET]
        values = {row[0].value: row[1].value for row in ws.iter_rows(min_row=1)}
        assert values["Зарплата User1"] == 120000
        assert values["Капитал"] == 500000  # unchanged
    finally:
        os.unlink(path)


def test_update_setting_raises_keyerror_if_missing():
    path = _make_workbook_with_settings({"Зарплата User1": 100000})
    try:
        writer = ExcelWriter(path)
        with pytest.raises(KeyError):
            writer.update_setting("Капитал", 700000)
    finally:
        os.unlink(path)


def test_add_mandatory_payment_writes_into_first_half_reserved_row(budget_path):
    writer = ExcelWriter(budget_path)
    row_num = writer.add_mandatory_payment(
        description="Парковка", amount=5000.0, due_day=6, half="first",
    )
    assert row_num > 0

    reader = ExcelReader(budget_path)
    payments = reader.get_mandatory_payments()
    parking = next(p for p in payments["first"] if p["description"] == "Парковка")
    assert parking["amount"] == 5000.0
    assert parking["due_day"] == 6
    assert all(p["description"] != "Парковка" for p in payments["second"])


def test_add_mandatory_payment_writes_into_second_half(budget_path):
    writer = ExcelWriter(budget_path)
    writer.add_mandatory_payment(
        description="Подписка", amount=500.0, due_day=0, half="second",
    )

    reader = ExcelReader(budget_path)
    payments = reader.get_mandatory_payments()
    subscription = next(p for p in payments["second"] if p["description"] == "Подписка")
    assert subscription["amount"] == 500.0
    # due_day=0 must be written as an empty "~" cell, not a literal 0.
    assert subscription["due_day"] == 0
    assert all(p["description"] != "Подписка" for p in payments["first"])


def test_add_mandatory_payment_does_not_touch_itogo_formula(budget_path):
    writer = ExcelWriter(budget_path)
    writer.add_mandatory_payment(
        description="Такси", amount=3000.0, due_day=8, half="first",
    )
    wb = openpyxl.load_workbook(budget_path)
    ws = wb["📅 Платежи"]
    itogo_row = next(
        r for r in range(2, ws.max_row + 1)
        if ws.cell(r, 1).value and "ИТОГО" in str(ws.cell(r, 1).value)
    )
    assert ws.cell(itogo_row, 2).value == "=SUM(C4:C17)"
    assert ws.cell(itogo_row, 3).value == "=SUM(C4:C17)"


def test_add_mandatory_payment_raises_when_block_is_full():
    wb = openpyxl.Workbook()
    ws_t = wb.active
    ws_t.title = "📋 Транзакции"
    ws_t.append([""] * 11)
    ws_t.append(["Дата"] + [""] * 10)

    ws_pay = wb.create_sheet("📅 Платежи")
    ws_pay.cell(2, 1).value = "Платёж / назначение"
    ws_pay.cell(3, 1).value = "Existing"
    ws_pay.cell(3, 2).value = 5
    ws_pay.cell(3, 3).value = 1000
    ws_pay.cell(4, 1).value = "ИТОГО"
    ws_pay.cell(4, 2).value = "=SUM(C3:C3)"
    ws_pay.cell(5, 1).value = "2️⃣  ВТОРАЯ ЗАРПЛАТА"
    ws_pay.cell(6, 1).value = "ИТОГО"

    path = tempfile.mktemp(suffix=".xlsx")
    wb.save(path)
    try:
        writer = ExcelWriter(path)
        with pytest.raises(ValueError, match="нет свободных строк"):
            writer.add_mandatory_payment(
                description="Новый платёж", amount=100.0, due_day=0, half="first",
            )
    finally:
        os.unlink(path)
