import shutil
from pathlib import Path

import openpyxl
import pytest

from excel.setup import (
    MandatoryPaymentSetup,
    WorkbookSetupData,
    apply_workbook_setup,
    validate_template_workbook,
)


SAMPLE_FILE = Path("example/budget_example.xlsx")


def test_validate_template_workbook_accepts_public_template():
    validate_template_workbook(SAMPLE_FILE)


def test_validate_template_workbook_rejects_missing_sheet(tmp_path):
    path = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "📋 Транзакции"
    wb.save(path)

    with pytest.raises(ValueError, match="missing required sheets"):
        validate_template_workbook(path)


def test_validate_template_workbook_rejects_private_values(tmp_path):
    path = tmp_path / "private.xlsx"
    shutil.copy(SAMPLE_FILE, path)
    wb = openpyxl.load_workbook(path)
    ws = wb["⚙️ Настройки"]
    ws["B4"] = "person@example.com"
    wb.save(path)

    with pytest.raises(ValueError, match="suspicious private"):
        validate_template_workbook(path)


def test_apply_workbook_setup_writes_runtime_only(tmp_path):
    template = tmp_path / "template.xlsx"
    runtime = tmp_path / "data" / "budget.xlsx"
    shutil.copy(SAMPLE_FILE, template)

    apply_workbook_setup(
        runtime,
        WorkbookSetupData(
            user_name="Tester",
            default_account="Debit",
            initial_capital=123000,
            salary=90000,
            salary_payment_model="fixed_percent",
            salary_payment_percentages=(0.4, 0.6),
            categories=["Food", "Transport"],
            accounts=["Debit", "Cash"],
            mandatory_payments=[
                MandatoryPaymentSetup("Rent", 30000, 5),
                MandatoryPaymentSetup("Phone", 1000, 20),
            ],
        ),
        template,
    )

    assert runtime.exists()
    original = openpyxl.load_workbook(template, data_only=True)
    assert original["⚙️ Настройки"]["B4"].value == "User1"

    wb = openpyxl.load_workbook(runtime, data_only=True)
    settings = {
        row[0]: row[1]
        for row in wb["⚙️ Настройки"].iter_rows(min_row=1, values_only=True)
        if row[0] is not None
    }
    assert settings["Имя 1"] == "Tester"
    assert settings["Зарплата Tester"] == 90000
    assert settings["salary_payment_model"] == "fixed_percent"
    assert wb["📈 Капитал"].cell(3, 2).value == 123000
    payments = [row[0] for row in wb["📅 Платежи"].iter_rows(min_row=4, values_only=True) if row[0]]
    assert "Rent" in payments
    assert "Phone" in payments
